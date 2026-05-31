#!/usr/bin/env python3
"""Export MRV report to Excel."""
import json
import os
import re
import sys
from collections import OrderedDict

LABELS_3 = ['entailment', 'neutral', 'contradiction']
LABELS_2 = ['entailment', 'not_entailment']

KNOWN_MRS = [
    'adding_contradiction', 'antonym_substitution', 'conditional_clause',
    'negation_flip', 'pronoun_substitution', 'summary',
    'synonym_replacement', 'uninformative', 'voice_switch'
]


def normalize_prediction_3(text):
    if not isinstance(text, str):
        return None
    t = text.strip().lower().rstrip('.!?')
    if t in LABELS_3:
        return t
    found = [lab for lab in LABELS_3 if lab in t]
    if found:
        return found[0]
    mapping = {'entailment': 'entailment', 'entailed': 'entailment', 'e': 'entailment',
               '0': 'entailment', 'neutral': 'neutral', 'n': 'neutral', '1': 'neutral',
               'contradiction': 'contradiction', 'contradictory': 'contradiction',
               'c': 'contradiction', 'contradict': 'contradiction', '2': 'contradiction'}
    t_clean = re.sub(r'[^a-z0-9]', '', t)
    return mapping.get(t_clean, None)


def normalize_prediction_binary(text):
    if not isinstance(text, str):
        return None
    t = text.strip().lower().rstrip('.!?')
    if t in ('entailment', 'not_entailment'):
        return t
    if t == 'entailment' or t in ('entailed', 'e', '0'):
        return 'entailment'
    for lab in ('not_entailment', 'neutral', 'contradiction', 'contradictory', 'contradict'):
        if lab in t:
            return 'not_entailment'
    mapping = {'entailment': 'entailment', 'entailed': 'entailment', 'e': 'entailment',
               '0': 'entailment', 'neutral': 'not_entailment', 'n': 'not_entailment',
               '1': 'not_entailment', 'contradiction': 'not_entailment',
               'contradictory': 'not_entailment', 'c': 'not_entailment',
               'contradict': 'not_entailment', '2': 'not_entailment',
               'not_entailment': 'not_entailment', 'not entailment': 'not_entailment'}
    t_clean = re.sub(r'[^a-z0-9]', '', t).replace(' ', '_')
    if t_clean in mapping:
        return mapping[t_clean]
    if 'not' in t and 'entail' in t:
        return 'not_entailment'
    if 'entail' in t:
        return 'entailment'
    return None


def extract_prediction(entry):
    resp = entry.get('response')
    pred_text = None
    if isinstance(resp, dict):
        choices = resp.get('choices', [])
        if choices:
            pred_text = choices[0].get('message', {}).get('content', '')
    elif isinstance(resp, str):
        pred_text = resp
    else:
        return None
    task = entry.get('meta', {}).get('task', 'nli')
    if task == 'nli-binary':
        return normalize_prediction_binary(pred_text)
    return normalize_prediction_3(pred_text)


def compute_mrv(base_dir):
    results = OrderedDict()
    dataset_totals = OrderedDict()

    datasets = sorted([d for d in os.listdir(base_dir)
                       if os.path.isdir(os.path.join(base_dir, d))])

    for ds in datasets:
        ds_dir = os.path.join(base_dir, ds)
        mr_files = sorted([f for f in os.listdir(ds_dir)
                          if f.endswith('.jsonl') and not f.endswith('_all.jsonl')])

        ds_total = 0
        ds_errors = 0
        ds_mrs = OrderedDict()

        for mr_file in mr_files:
            mr = mr_file.replace('.jsonl', '')
            mr_path = os.path.join(ds_dir, mr_file)
            total = 0
            errors = 0

            with open(mr_path) as f:
                for line in f:
                    line = line.strip()
                    if not line:
                        continue
                    try:
                        entry = json.loads(line)
                    except json.JSONDecodeError:
                        continue
                    if entry.get('status') != 200:
                        continue
                    gold = entry.get('meta', {}).get('gold_label')
                    if gold is None:
                        continue
                    pred = extract_prediction(entry)
                    if pred is None:
                        continue
                    total += 1
                    if gold != pred:
                        errors += 1

            if total > 0:
                ds_mrs[mr] = {'total': total, 'errors': errors, 'mrv': errors / total}
                ds_total += total
                ds_errors += errors

        results[ds] = ds_mrs
        dataset_totals[ds] = {'total': ds_total, 'errors': ds_errors,
                              'mrv': ds_errors / ds_total if ds_total > 0 else 0}

    return results, dataset_totals


def to_excel(results, dataset_totals, output_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # ===== Sheet 1: MRV Matrix =====
    ws1 = wb.active
    ws1.title = 'MRV汇总'

    # Collect all MRs that have data, preserving order from KNOWN_MRS
    seen = set()
    ordered_mrs = []
    for mr in KNOWN_MRS:
        for ds_mrs in results.values():
            if mr in ds_mrs and mr not in seen:
                ordered_mrs.append(mr)
                seen.add(mr)
                break

    # ── styles ──
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill('solid', fgColor='4472C4')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell_align = Alignment(horizontal='center', vertical='center')
    cell_align_left = Alignment(horizontal='left', vertical='center')
    bold_font = Font(bold=True, size=11)
    bold_font_blue = Font(bold=True, color='FFFFFF', size=11)
    blue_fill = PatternFill('solid', fgColor='D6E4F0')
    light_fill = PatternFill('solid', fgColor='F2F2F2')
    thin_border = Border(
        left=Side(style='thin', color='B4B4B4'),
        right=Side(style='thin', color='B4B4B4'),
        top=Side(style='thin', color='B4B4B4'),
        bottom=Side(style='thin', color='B4B4B4')
    )
    green_fill = PatternFill('solid', fgColor='C6EFCE')
    yellow_fill = PatternFill('solid', fgColor='FFEB9C')
    red_fill = PatternFill('solid', fgColor='FFC7CE')

    # ── headers ──
    row = 1
    headers = ['数据集', '总数', '错误数', 'MRV'] + ordered_mrs
    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # ── data rows ──
    for i, ds in enumerate(results):
        row += 1
        dt = dataset_totals[ds]
        values = [ds, dt['total'], dt['errors'], round(dt['mrv'], 4)]
        for mr in ordered_mrs:
            if mr in results[ds]:
                r = results[ds][mr]
                values.append(round(r['mrv'], 4))
            else:
                values.append(None)

        for col_idx, v in enumerate(values, 1):
            cell = ws1.cell(row=row, column=col_idx, value=v)
            cell.alignment = cell_align
            cell.border = thin_border
            if i % 2 == 0:
                cell.fill = light_fill

            # Color-code MRV cells
            if col_idx > 4 and v is not None:
                if v <= 0.1:
                    cell.fill = green_fill
                elif v <= 0.2:
                    cell.fill = yellow_fill
                else:
                    cell.fill = red_fill

        # Dataset name left-aligned, bold
        ws1.cell(row=row, column=1).alignment = cell_align_left
        ws1.cell(row=row, column=1).font = Font(bold=True)

    # ── ALL row ──
    row += 1
    grand_total = sum(dt['total'] for dt in dataset_totals.values())
    grand_errors = sum(dt['errors'] for dt in dataset_totals.values())
    grand_mrv = round(grand_errors / grand_total, 4) if grand_total > 0 else 0
    all_values = ['ALL', grand_total, grand_errors, grand_mrv]
    for mr in ordered_mrs:
        # Weighted average MRV for "ALL" per MR
        mr_total = sum(results[ds][mr]['total'] for ds in results if mr in results[ds])
        mr_errors = sum(results[ds][mr]['errors'] for ds in results if mr in results[ds])
        all_values.append(round(mr_errors / mr_total, 4) if mr_total > 0 else None)

    for col_idx, v in enumerate(all_values, 1):
        cell = ws1.cell(row=row, column=col_idx, value=v)
        cell.font = Font(bold=True, size=11)
        cell.alignment = cell_align
        cell.border = Border(
            left=Side(style='thin', color='4472C4'),
            right=Side(style='thin', color='4472C4'),
            top=Side(style='medium', color='4472C4'),
            bottom=Side(style='medium', color='4472C4')
        )
        cell.fill = PatternFill('solid', fgColor='B4C6E7')

    ws1.cell(row=row, column=1).alignment = cell_align_left

    # ── column widths ──
    ws1.column_dimensions['A'].width = 10
    ws1.column_dimensions['B'].width = 8
    ws1.column_dimensions['C'].width = 8
    ws1.column_dimensions['D'].width = 8
    for j, mr in enumerate(ordered_mrs):
        ws1.column_dimensions[get_column_letter(5 + j)].width = max(14, len(mr) + 2)

    # ===== Sheet 2: Detail =====
    ws2 = wb.create_sheet('详细数据')

    row = 1
    detail_headers = ['数据集', 'MR类型', '总用例数', '错误数', '正确数', 'MRV (错误率)', '准确率']
    for col_idx, h in enumerate(detail_headers, 1):
        cell = ws2.cell(row=row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    row = 1
    for ds in results:
        for mr in ordered_mrs:
            if mr not in results[ds]:
                continue
            row += 1
            r = results[ds][mr]
            correct = r['total'] - r['errors']
            accuracy = correct / r['total']
            vals = [ds, mr, r['total'], r['errors'], correct, round(r['mrv'], 4), round(accuracy, 4)]
            for col_idx, v in enumerate(vals, 1):
                cell = ws2.cell(row=row, column=col_idx, value=v)
                cell.alignment = cell_align if col_idx > 2 else cell_align_left
                cell.border = thin_border
                if row % 2 == 0:
                    cell.fill = light_fill
                # Color MRV column
                if col_idx == 6:
                    if v is not None and v <= 0.1:
                        cell.fill = green_fill
                    elif v is not None and v <= 0.2:
                        cell.fill = yellow_fill
                    elif v is not None:
                        cell.fill = red_fill

            ws2.cell(row=row, column=1).font = Font(bold=True)
            ws2.cell(row=row, column=2).font = Font(bold=True)

    # Column widths
    ws2.column_dimensions['A'].width = 10
    ws2.column_dimensions['B'].width = 24
    for c in 'CDEFG':
        ws2.column_dimensions[c].width = 14

    wb.save(output_path)
    print(f'Saved: {output_path}')


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='Export MRV report to Excel')
    parser.add_argument('input_dir', nargs='?', default='output',
                        help='Model output directory (default: output)')
    parser.add_argument('--output', '-o', default=None,
                        help='Output .xlsx path')
    args = parser.parse_args()

    base = args.input_dir
    if not os.path.isdir(base):
        print(f'Directory not found: {base}')
        sys.exit(1)

    results, dataset_totals = compute_mrv(base)

    output_path = args.output or os.path.join(base, 'mrv_report.xlsx')
    to_excel(results, dataset_totals, output_path)
