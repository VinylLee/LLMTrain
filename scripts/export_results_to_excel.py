#!/usr/bin/env python3
"""Export all experiment results to Excel with detailed columns."""
import json
from pathlib import Path
from collections import defaultdict

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]

def compute_accuracy(jsonl_path):
    if not jsonl_path.exists():
        return None, 0
    correct = total = 0
    with open(jsonl_path) as f:
        for line in f:
            d = json.loads(line)
            if d.get("correct"):
                correct += 1
            total += 1
    return correct, total

def parse_experiment_label(exp_name):
    """Extract structured info from experiment name."""
    parts = exp_name.split("_")

    # Find seed
    seed = None
    for p in parts:
        if p.startswith("seed") and p[4:].isdigit():
            seed = p[4:]

    # MR-as-Instruction mode
    mode_map = {
        "none": "none", "pairop": "pair_operation", "pair_operation": "pair_operation",
        "shuffled": "shuffled_operation", "shuffled_operation": "shuffled_operation",
        "mrop": "operation_only", "operation_only": "operation_only",
        "paironly": "pair_only", "pair_only": "pair_only",
        "fulloracle": "full_oracle", "full_oracle": "full_oracle",
    }
    for p in parts:
        if p in mode_map:
            ds_name_map = {"mnlim": "MNLI-m", "sick": "SICK", "snli": "SNLI"}
            train_label = "MetTrain augmented"
            for i, pp in enumerate(parts):
                if pp in ds_name_map:
                    ds_name = ds_name_map[pp]
                    target = parts[i+1] if i+1 < len(parts) and parts[i+1].isdigit() else "?"
                    train_label = f"{ds_name} {target} (MetTrain augmented)"
                    break
            return {
                "experiment_type": "MR-as-Instruction",
                "mode": mode_map[p],
                "train_data": train_label,
                "seed": seed,
            }

    # Phase B: standard MetTrain
    ds_name_map = {"mnlim": "MNLI-m", "mnlimm": "MNLI-mm", "sick": "SICK", "snli": "SNLI"}

    if "original" in exp_name:
        exp_type = "Original Fine-tune"
        for i, p in enumerate(parts):
            if p in ds_name_map:
                ds_name = ds_name_map[p]
                target = parts[i+1] if i+1 < len(parts) and parts[i+1].isdigit() else "?"
                train_label = f"{ds_name} {target} (Original)"
                return {
                    "experiment_type": exp_type,
                    "mode": "N/A (standard NLI)",
                    "train_data": train_label,
                    "seed": seed,
                }

    if "mettrain" in exp_name:
        exp_type = "MetTrain Fine-tune"
        for i, p in enumerate(parts):
            if p in ds_name_map:
                ds_name = ds_name_map[p]
                target = parts[i+1] if i+1 < len(parts) and parts[i+1].isdigit() else "?"
                train_label = f"{ds_name} {target} (MetTrain augmented)"
                return {
                    "experiment_type": exp_type,
                    "mode": "N/A (standard NLI)",
                    "train_data": train_label,
                    "seed": seed,
                }

    return {"experiment_type": "Unknown", "mode": "?", "train_data": exp_name, "seed": seed}

def collect_all_results():
    rows = []
    exp_base = ROOT / "output" / "experiments"

    # Collect all roots: subdirectories with _progress.json + the base dir itself
    roots_to_scan = []
    for d in sorted(exp_base.iterdir()):
        if d.is_dir() and (d / "_progress.json").exists():
            roots_to_scan.append(d)
    # Phase B lives directly under exp_base
    if (exp_base / "_progress.json").exists():
        roots_to_scan.append(exp_base)

    for exp_root in roots_to_scan:
        progress_file = exp_root / "_progress.json"
        progress = json.loads(progress_file.read_text(encoding="utf-8"))

        for exp_name in sorted(progress):
            exp_dir = exp_root / exp_name
            if not exp_dir.is_dir():
                continue
            info = parse_experiment_label(exp_name)
            if not info["seed"]:
                continue

            # Original tests
            orig_dir = exp_dir / "tests" / "original"
            if orig_dir.is_dir():
                for f in sorted(orig_dir.glob("*.jsonl")):
                    correct, total = compute_accuracy(f)
                    if total > 0:
                        rows.append({
                            "实验组": exp_root.name,
                            "实验名称": exp_name,
                            "实验类型": info["experiment_type"],
                            "MR模式": info["mode"],
                            "训练数据": info["train_data"],
                            "种子": info["seed"],
                            "测试类型": "Original",
                            "测试集": f.stem.upper(),
                            "正确数": correct,
                            "总数": total,
                            "准确率(%)": round(correct / total * 100, 2),
                        })

            # MR tests
            mr_dir = exp_dir / "tests" / "mr"
            if mr_dir.is_dir():
                for f in sorted(mr_dir.glob("*.jsonl")):
                    correct, total = compute_accuracy(f)
                    if total > 0:
                        rows.append({
                            "实验组": exp_root.name,
                            "实验名称": exp_name,
                            "实验类型": info["experiment_type"],
                            "MR模式": info["mode"],
                            "训练数据": info["train_data"],
                            "种子": info["seed"],
                            "测试类型": "MR",
                            "测试集": f.stem.upper(),
                            "正确数": correct,
                            "总数": total,
                            "准确率(%)": round(correct / total * 100, 2),
                        })
    return rows

def write_excel(rows, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验结果汇总"

    headers = ["实验组", "实验名称", "实验类型", "MR模式", "训练数据", "种子",
               "测试类型", "测试集", "正确数", "总数", "准确率(%)"]

    # Styles
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    # Write header
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Write data
    for r, row in enumerate(rows, 2):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=r, column=col, value=row.get(h, ""))
            cell.border = thin_border
            cell.alignment = data_align
            # Highlight accuracy >= 90%
            if h == "准确率(%)" and isinstance(row.get(h), (int, float)) and row[h] >= 90:
                cell.font = Font(bold=True, color="006100")
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            # Highlight accuracy < 70%
            elif h == "准确率(%)" and isinstance(row.get(h), (int, float)) and row[h] < 70:
                cell.font = Font(bold=True, color="9C0006")
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    # Column widths
    col_widths = [40, 55, 20, 18, 35, 8, 12, 10, 10, 10, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freeze header
    ws.freeze_panes = "A2"

    # Auto filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"

    # Summary sheet
    ws2 = wb.create_sheet("汇总统计")
    ws2.append(["Phase", "实验类型", "种子数", "实验数", "测试集数", "总记录数"])

    # Phase A summary
    for exp_group in set(r["实验组"] for r in rows if "mrinstr" in r["实验组"].lower()):
        group_rows = [r for r in rows if r["实验组"] == exp_group]
        ws2.append([
            "Phase A: MR-as-Instruction",
            exp_group,
            len(set(r["种子"] for r in group_rows)),
            len(set(r["实验名称"] for r in group_rows)),
            len(set((r["测试类型"], r["测试集"]) for r in group_rows)),
            len(group_rows),
        ])

    # Phase B summary
    phase_b = [r for r in rows if r["实验类型"] != "MR-as-Instruction"]
    ws2.append([
        "Phase B: MetTrain",
        "experiments/configs/experiments_config.json",
        len(set(r["种子"] for r in phase_b)),
        len(set(r["实验名称"] for r in phase_b)),
        len(set((r["测试类型"], r["测试集"]) for r in phase_b)),
        len(phase_b),
    ])

    ws2.append([])
    ws2.append([f"总记录数: {len(rows)}"])
    ws2.append([f"输出文件: {output_path}"])

    wb.save(output_path)
    return len(rows)

if __name__ == "__main__":
    print("收集实验结果...")
    rows = collect_all_results()
    print(f"共 {len(rows)} 条记录")

    output = ROOT / "artifacts" / "experiment_results_full.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    count = write_excel(rows, output)
    print(f"已保存: {output}")
    print(f"总计 {count} 条结果")
